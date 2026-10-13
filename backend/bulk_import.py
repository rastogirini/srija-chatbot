"""Bulk product import for the admin "Bulk Upload" page - a shopkeeper
receiving a whole shipment fills one spreadsheet row per product (CSV or
Excel) and selects all the product photos at once; each row's `image_file`
column names the photo that belongs to it.

The flow is two calls with the same files: a preview (validate every row,
save nothing) so the shopkeeper can see problems before anything goes live,
then a commit that re-validates server-side and saves the good rows in one
transaction. Bad rows are skipped rather than blocking the whole upload.

The uploaded spreadsheet itself is never stored - only an import_batches
summary row, and each product it created tagged with that batch's id. An
old upload can then be re-downloaded by rebuilding it from those product
rows (see export_products_xlsx), showing current values, not stale ones."""

import csv
import io
import os
import re
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from psycopg2.extras import RealDictCursor

from db import get_db_connection
from products import UPLOAD_DIR, _allowed_image, insert_product, save_uploaded_image

TEMPLATE_COLUMNS = [
    "category", "name", "price", "description", "subcategory",
    "cuisine", "diet", "course", "stock", "image_file",
]
REQUIRED_COLUMNS = {"category", "name", "price"}
MAX_ROWS = 500

CATEGORIES = ("food", "beauty", "toys")
# Same option lists as the single "Add Product" form (AddProduct.jsx), so a
# bulk-imported product is filterable on the main site exactly like a
# hand-added one. Diet also allows the calorie tags the migrated menu uses.
CUISINES = ("italian", "indian", "chinese")
COURSES = ("appetizer", "main", "dessert", "bread", "salad")
DIET_TAGS = ("vegetarian", "vegan", "gluten-free", "low-calorie", "high-calorie")
SUBCATEGORIES = {
    "beauty": ("lipstick", "foundation", "skincare", "fragrance"),
    "toys": ("building-blocks", "action-figures", "board-games", "puzzles", "outdoor"),
}

TEMPLATE_EXAMPLES = [
    ["toys", "LEGO City Fire Truck", 2499, "214-piece fire truck set with 2 minifigures", "building-blocks", "", "", "", 10, "lego_fire.jpg"],
    ["beauty", "Matte Lipstick - Ruby Red", 499, "Long-lasting matte finish", "lipstick", "", "", "", 25, "lipstick_red.png"],
    ["food", "Paneer Tikka", 320, "Chargrilled cottage cheese with mint chutney", "", "indian", "vegetarian, gluten-free", "appetizer", "", "paneer.jpg"],
]


class SpreadsheetError(Exception):
    """The file as a whole can't be read (wrong type, missing required
    columns, too many rows) - as opposed to a single bad row, which is
    reported per row and doesn't stop the others."""


def _normalize_header(value):
    return re.sub(r"\s+", "_", str(value or "").strip().lower())


def _cell_to_str(value):
    """Excel hands back numbers as int/float (2499 -> 2499.0) - turn whole
    floats back into "2499" so they read naturally in the preview and parse
    cleanly as stock counts."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_spreadsheet(file_storage):
    """Turn an uploaded .csv/.xlsx into a list of {column: str} dicts, one
    per non-blank row, each tagged with `_row` - its row number as the
    shopkeeper sees it in Excel (header is row 1), so error messages point
    at the right line. Raises SpreadsheetError for whole-file problems."""
    if not file_storage or not file_storage.filename:
        raise SpreadsheetError("Please choose a spreadsheet file.")

    filename = file_storage.filename.lower()
    raw = file_storage.read()

    if filename.endswith(".csv"):
        # Excel's "Save as CSV" writes a UTF-8 BOM; older Windows Excel
        # writes cp1252 instead - accept both rather than failing on ₹.
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252", errors="replace")
        table = list(csv.reader(io.StringIO(text)))
    elif filename.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise SpreadsheetError("That Excel file couldn't be opened - is it a valid .xlsx file?")
        sheet = workbook.worksheets[0]
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    else:
        raise SpreadsheetError("Only .csv and .xlsx files are supported.")

    if not table:
        raise SpreadsheetError("The spreadsheet is empty.")

    headers = [_normalize_header(h) for h in table[0]]
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise SpreadsheetError(
            f"Missing required column(s): {', '.join(sorted(missing))}. "
            "Download the template to see the expected columns."
        )

    rows = []
    for index, values in enumerate(table[1:], start=2):
        cells = {header: _cell_to_str(value) for header, value in zip(headers, values) if header}
        if not any(cells.values()):
            continue
        cells["_row"] = index
        rows.append(cells)

    if not rows:
        raise SpreadsheetError("The spreadsheet has a header row but no products under it.")
    if len(rows) > MAX_ROWS:
        raise SpreadsheetError(f"Too many rows ({len(rows)}) - please split it into files of at most {MAX_ROWS} products.")

    return rows


def _existing_product_names():
    """Lowercased current and former names of every live product, for
    duplicate detection - same rule as products.name_taken."""
    conn = get_db_connection()
    if not conn:
        return None
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT LOWER(known_name)
            FROM products, unnest(array_prepend(name, COALESCE(former_names, '{}'))) AS known_name
            WHERE is_deleted = FALSE
        """)
        return {row[0] for row in cur.fetchall()}
    except Exception as e:
        print(f"❌ Error reading product names: {e}")
        return None
    finally:
        conn.close()


def _parse_price(raw):
    """Accept the ways a shopkeeper actually types a price: 2499, 2,499,
    ₹2499, Rs. 2499."""
    cleaned = re.sub(r"(?i)^\s*(rs\.?|inr|₹)\s*", "", raw).replace(",", "").strip()
    try:
        price = float(cleaned)
    except ValueError:
        return None
    return price if price > 0 else None


def validate_rows(rows, image_names):
    """Check every row and return one result per row:
    {row, status, messages, data, raw}. `status` is:
      - "ok"        - will be imported as-is
      - "warning"   - will be imported, but something's worth a look
                      (no matching photo, an ignored column)
      - "duplicate" - skipped: a live product already has this name
      - "error"     - skipped: the row can't be saved as written
    `image_names` is the lowercased set of uploaded photo filenames.
    Returns None if the database can't be reached for the duplicate check."""
    existing_names = _existing_product_names()
    if existing_names is None:
        return None

    seen_in_file = {}
    results = []

    for row in rows:
        errors, warnings = [], []
        category = row.get("category", "").lower()
        name = row.get("name", "")

        if category not in CATEGORIES:
            errors.append(f"category must be one of: {', '.join(CATEGORIES)}" if category else "category is missing")
        if not name:
            errors.append("name is missing")
        elif len(name) > 200:
            errors.append("name is longer than 200 characters")

        price = None
        if not row.get("price"):
            errors.append("price is missing")
        else:
            price = _parse_price(row["price"])
            if price is None:
                errors.append(f'price "{row["price"]}" is not a positive number')

        stock = None
        if row.get("stock"):
            try:
                stock = int(row["stock"])
                if stock < 0:
                    errors.append("stock cannot be negative")
            except ValueError:
                errors.append(f'stock "{row["stock"]}" is not a whole number (leave it blank for unlimited)')

        cuisine = course = diet = subcategory = None
        if category == "food":
            cuisine = row.get("cuisine", "").lower() or None
            if cuisine and cuisine not in CUISINES:
                errors.append(f'cuisine "{cuisine}" must be one of: {", ".join(CUISINES)}')
            course = row.get("course", "").lower() or None
            if course and course not in COURSES:
                errors.append(f'course "{course}" must be one of: {", ".join(COURSES)}')
            diet_tags = [t.strip().lower() for t in re.split(r"[,;]", row.get("diet", "")) if t.strip()]
            bad_tags = [t for t in diet_tags if t not in DIET_TAGS]
            if bad_tags:
                errors.append(f'diet "{", ".join(bad_tags)}" must be from: {", ".join(DIET_TAGS)}')
            diet = ",".join(dict.fromkeys(diet_tags)) or None
            if row.get("subcategory"):
                warnings.append("subcategory is only used for beauty/toys - ignored for food")
        elif category in SUBCATEGORIES:
            subcategory = row.get("subcategory", "").lower() or None
            allowed = SUBCATEGORIES[category]
            if subcategory and subcategory not in allowed:
                errors.append(f'subcategory "{subcategory}" must be one of: {", ".join(allowed)}')
            if any(row.get(col) for col in ("cuisine", "diet", "course")):
                warnings.append("cuisine/diet/course are only used for food - ignored")

        image_file = os.path.basename(row.get("image_file", "")).lower() or None
        if not image_file:
            warnings.append("no image_file given - will be saved without a photo")
        elif not _allowed_image(image_file):
            warnings.append(f'"{image_file}" is not a supported image type - will be saved without a photo')
            image_file = None
        elif image_file not in image_names:
            warnings.append(f'photo "{image_file}" was not among the uploaded photos - will be saved without a photo')
            image_file = None

        status = "error" if errors else "warning" if warnings else "ok"
        lowered = name.lower()
        if not errors and lowered in existing_names:
            status = "duplicate"
            errors.append("a product with this name already exists - skipped")
        elif not errors and lowered in seen_in_file:
            status = "error"
            errors.append(f"same name as row {seen_in_file[lowered]} in this file - skipped")
        if name and lowered not in seen_in_file:
            seen_in_file[lowered] = row["_row"]

        results.append({
            "row": row["_row"],
            "status": status,
            # A skipped row only needs the reason it was skipped - "no
            # photo" noise on top of that would just bury it.
            "messages": errors if errors else warnings,
            "image_file": image_file,
            "raw": {col: row.get(col, "") for col in TEMPLATE_COLUMNS},
            "data": None if status in ("error", "duplicate") else {
                "category": category,
                "name": name,
                "price": price,
                "description": row.get("description", ""),
                "cuisine": cuisine,
                "diet": diet,
                "course": course,
                "subcategory": subcategory,
                "stock": stock,
            },
        })

    return results


def summarize(results):
    return {
        "ok": sum(1 for r in results if r["status"] == "ok"),
        "warning": sum(1 for r in results if r["status"] == "warning"),
        "duplicate": sum(1 for r in results if r["status"] == "duplicate"),
        "error": sum(1 for r in results if r["status"] == "error"),
    }


def import_rows(filename, results, images_by_name):
    """Save every importable row (ok/warning) plus its photo, all in one
    transaction under a new import_batches row. If anything fails part-way,
    the whole batch rolls back and any photos already written are removed,
    so a failed upload never leaves half a shipment live. Returns
    (batch_id, added_count), or None on failure."""
    importable = [r for r in results if r["data"] is not None]
    counts = summarize(results)

    conn = get_db_connection()
    if not conn:
        return None

    saved_images = []
    try:
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO import_batches (filename, added_count, skipped_count, warning_count)
            VALUES (%s, %s, %s, %s) RETURNING id
        """, (filename, len(importable), counts["error"] + counts["duplicate"], counts["warning"]))
        batch_id = cur.fetchone()[0]

        for result in importable:
            image_url = None
            image = images_by_name.get(result["image_file"]) if result["image_file"] else None
            if image:
                # Two rows may share one photo - rewind so the second save
                # doesn't write an empty file from an already-read stream.
                image.stream.seek(0)
                image_url = save_uploaded_image(image)
                if image_url:
                    saved_images.append(image_url)
            insert_product(cur, result["data"], image_url, import_batch_id=batch_id)

        conn.commit()
        cur.close()
        print(f"✅ Bulk import #{batch_id}: {len(importable)} products from {filename}")
        return batch_id, len(importable)
    except Exception as e:
        conn.rollback()
        print(f"❌ Error during bulk import: {e}")
        for image_url in saved_images:
            path = os.path.join(UPLOAD_DIR, os.path.basename(image_url))
            try:
                if os.path.exists(path):
                    os.remove(path)
            except OSError:
                pass
        return None
    finally:
        conn.close()


def list_import_batches():
    """Every past bulk upload, newest first, with how many of its products
    are still live (some may have been deleted since)."""
    conn = get_db_connection()
    if not conn:
        return None
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("""
            SELECT b.*,
                   (SELECT COUNT(*) FROM products p
                    WHERE p.import_batch_id = b.id AND p.is_deleted = FALSE) AS live_count
            FROM import_batches b
            ORDER BY b.created_at DESC
        """)
        rows = cur.fetchall()
        cur.close()
        return rows
    except Exception as e:
        print(f"❌ Error listing import batches: {e}")
        return None
    finally:
        conn.close()


def get_batch_products(batch_id):
    """(batch, products) for one upload - products include deleted ones,
    since "what came in this shipment" should still list an item that was
    later pulled; the export marks them. Returns (None, None) if not found."""
    conn = get_db_connection()
    if not conn:
        return None, None
    try:
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("SELECT * FROM import_batches WHERE id = %s", (batch_id,))
        batch = cur.fetchone()
        if not batch:
            return None, None
        cur.execute("SELECT * FROM products WHERE import_batch_id = %s ORDER BY id", (batch_id,))
        products = cur.fetchall()
        cur.close()
        return batch, products
    except Exception as e:
        print(f"❌ Error reading import batch: {e}")
        return None, None
    finally:
        conn.close()


_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="667EEA")


def _style_header(sheet, widths):
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"


_COLUMN_WIDTHS = {"A": 10, "B": 32, "C": 10, "D": 44, "E": 18, "F": 11, "G": 26, "H": 12, "I": 9, "J": 22, "K": 34, "L": 10}


def _to_bytes(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def build_template_xlsx():
    """The blank template: the expected columns with three example rows
    (one per category) to copy the shape of, plus an "Allowed values"
    sheet so the shopkeeper doesn't have to guess the exact spellings."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(TEMPLATE_COLUMNS)
    for example in TEMPLATE_EXAMPLES:
        sheet.append(example)
    _style_header(sheet, _COLUMN_WIDTHS)

    help_sheet = workbook.create_sheet("Allowed values")
    help_sheet.append(["column", "allowed values", "notes"])
    help_rows = [
        ["category", ", ".join(CATEGORIES), "Required."],
        ["name", "", "Required. Must not match an existing product's name."],
        ["price", "", "Required, in ₹. 2499, 2,499 and ₹2499 all work."],
        ["description", "", "Optional."],
        ["subcategory", "beauty: " + ", ".join(SUBCATEGORIES["beauty"]), "Beauty and toys only."],
        ["", "toys: " + ", ".join(SUBCATEGORIES["toys"]), ""],
        ["cuisine", ", ".join(CUISINES), "Food only."],
        ["diet", ", ".join(DIET_TAGS), "Food only. Several allowed, comma-separated."],
        ["course", ", ".join(COURSES), "Food only."],
        ["stock", "", "Optional whole number. Leave blank for unlimited (untracked)."],
        ["image_file", "", "The photo's exact filename, e.g. lego_fire.jpg. Select all the photos together when uploading."],
    ]
    for row in help_rows:
        help_sheet.append(row)
    _style_header(help_sheet, {"A": 14, "B": 70, "C": 70})

    return _to_bytes(workbook)


def export_products_xlsx(products, base_url, title="Products"):
    """Products as a spreadsheet in the same column layout as the template
    (so it reads the same as what was uploaded), with current values.
    image_file is replaced by image_url (stored photos are renamed on
    upload, so the original filename no longer exists) - made absolute
    with `base_url` so it's clickable from Excel - plus a status column
    that marks products deleted since."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    columns = [c for c in TEMPLATE_COLUMNS if c != "image_file"] + ["image_url", "status"]
    sheet.append(columns)
    for product in products:
        row = []
        for column in columns:
            if column == "status":
                row.append("deleted" if product.get("is_deleted") else "live")
            elif column == "image_url":
                image = product.get("image") or ""
                row.append(base_url + image if image.startswith("/") else image)
            else:
                value = product.get(column)
                row.append(float(value) if isinstance(value, Decimal) else value)
        sheet.append(row)
    _style_header(sheet, _COLUMN_WIDTHS)
    return _to_bytes(workbook)
